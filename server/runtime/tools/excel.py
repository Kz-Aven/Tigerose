"""Excel (.xlsx) read/write tools — openpyxl backed, practical fallbacks."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any

from server.runtime.tools import files as file_tools

_MAX_OUT = 50000
_DEFAULT_MAX_ROWS = 200
_CELL_REF = re.compile(r"^\$?([A-Za-z]+)\$?(\d+)$")


def _openpyxl():
    try:
        import openpyxl  # type: ignore
    except ImportError:
        return None
    return openpyxl


def _missing_dep() -> str:
    return json.dumps(
        {
            "ok": False,
            "error": "openpyxl not installed; pip install openpyxl (or reinstall App runtime)",
        },
        ensure_ascii=False,
    )


def _is_xlsx(path: Path) -> bool:
    return path.suffix.lower() == ".xlsx"


def _reject_ext(path: Path) -> str | None:
    suf = path.suffix.lower()
    if suf == ".xlsx":
        return None
    if suf in (".xls", ".xlsm", ".xlsb"):
        return json.dumps(
            {
                "ok": False,
                "error": f"unsupported Excel format '{suf}'; v1 only supports .xlsx",
            },
            ensure_ascii=False,
        )
    return json.dumps(
        {
            "ok": False,
            "error": f"not an .xlsx file ({path.name}); use read_file/write_file for text formats",
        },
        ensure_ascii=False,
    )


def _col_row(cell_ref: str) -> tuple[int, int]:
    m = _CELL_REF.match(cell_ref.strip())
    if not m:
        raise ValueError(f"invalid cell ref: {cell_ref}")
    col_s, row_s = m.group(1).upper(), m.group(2)
    col = 0
    for ch in col_s:
        col = col * 26 + (ord(ch) - 64)
    return col, int(row_s)


def _parse_range(range_a1: str) -> tuple[int, int, int, int]:
    """Return (min_col, min_row, max_col, max_row) 1-based."""
    raw = range_a1.strip().replace("$", "")
    if ":" in raw:
        a, b = raw.split(":", 1)
        c1, r1 = _col_row(a)
        c2, r2 = _col_row(b)
        return min(c1, c2), min(r1, r2), max(c1, c2), max(r1, r2)
    c, r = _col_row(raw)
    return c, r, c, r


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _pick_sheet(wb, sheet: str | None):
    if sheet:
        if sheet not in wb.sheetnames:
            raise KeyError(f"sheet not found: {sheet}; available={wb.sheetnames}")
        return wb[sheet]
    # Prefer first visible
    for name in wb.sheetnames:
        ws = wb[name]
        state = getattr(ws, "sheet_state", "visible") or "visible"
        if state == "visible":
            return ws
    return wb[wb.sheetnames[0]]


def _used_bounds(ws) -> tuple[int, int, int, int]:
    dim = ws.dimensions
    if dim and dim != "A1:A1":
        try:
            return _parse_range(dim)
        except ValueError:
            pass
    if ws.max_row and ws.max_column:
        return 1, 1, int(ws.max_column), int(ws.max_row)
    return 1, 1, 1, 1


def _matrix_from_ws(
    ws,
    *,
    range_a1: str | None,
    max_rows: int,
) -> tuple[list[list[str]], str, bool]:
    truncated = False

    if range_a1:
        min_c, min_r, max_c, max_r = _parse_range(range_a1)
        used = f"{_a1(min_c, min_r)}:{_a1(max_c, max_r)}"
        if max_r - min_r + 1 > max_rows:
            max_r = min_r + max_rows - 1
            truncated = True
        rows: list[list[str]] = []
        for r in range(min_r, max_r + 1):
            row = [_cell_str(ws.cell(row=r, column=c).value) for c in range(min_c, max_c + 1)]
            rows.append(row)
        return rows, used, truncated

    # Prefer dimensions when available; otherwise stream via iter_rows.
    try:
        min_c, min_r, max_c, max_r = _used_bounds(ws)
        used = f"{_a1(min_c, min_r)}:{_a1(max_c, max_r)}"
        if max_r - min_r + 1 > max_rows:
            max_r = min_r + max_rows - 1
            truncated = True
            used = f"{_a1(min_c, min_r)}:{_a1(max_c, max_r)} (truncated)"
        rows = []
        for r in range(min_r, max_r + 1):
            row = [_cell_str(ws.cell(row=r, column=c).value) for c in range(min_c, max_c + 1)]
            rows.append(row)
        if rows:
            return rows, used, truncated
    except Exception:  # noqa: BLE001
        pass

    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= max_rows:
            truncated = True
            break
        rows.append([_cell_str(v) for v in row])
    used = f"streamed:{len(rows)}rows"
    return rows, used, truncated


def _a1(col: int, row: int) -> str:
    s = ""
    n = col
    while n:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return f"{s}{row}"


def _format_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return "(empty)"
    width = max(len(r) for r in rows)
    norm = [r + [""] * (width - len(r)) for r in rows]
    header = norm[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for r in norm[1:]:
        lines.append("| " + " | ".join(r) + " |")
    if len(norm) == 1:
        # only header-like single row — still show as one-row table without fake body
        return lines[0] + "\n" + lines[1]
    return "\n".join(lines)


def _format_csv(rows: list[list[str]]) -> str:
    import io

    buf = io.StringIO()
    w = csv.writer(buf)
    for r in rows:
        w.writerow(r)
    return buf.getvalue().rstrip("\n")


def excel_read(
    cwd: Path,
    path: str,
    *,
    sheet: str | None = None,
    range: str | None = None,  # noqa: A002 — tool API name
    max_rows: int | None = None,
    format: str = "markdown",  # noqa: A002
    trusted_roots: tuple[Path, ...] = (),
) -> str:
    ox = _openpyxl()
    if ox is None:
        return _missing_dep()

    try:
        target = file_tools.ensure_under(cwd, file_tools.resolve_path(cwd, path), trusted_roots)
    except PermissionError as exc:
        return json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False)

    bad = _reject_ext(target)
    if bad:
        return bad
    if not target.is_file():
        return json.dumps({"ok": False, "error": f"File not found: {path}"}, ensure_ascii=False)

    fmt = (format or "markdown").lower().strip()
    if fmt not in ("markdown", "csv"):
        return json.dumps({"ok": False, "error": "format must be markdown|csv"}, ensure_ascii=False)

    try:
        wb = ox.load_workbook(target, read_only=True, data_only=True)
        try:
            ws = _pick_sheet(wb, sheet)
            rows, used, truncated = _matrix_from_ws(
                ws,
                range_a1=range,
                max_rows=int(max_rows or _DEFAULT_MAX_ROWS),
            )
            body = _format_markdown(rows) if fmt == "markdown" else _format_csv(rows)
            header = (
                f"ok=true path={target} sheet={ws.title} used_range={used} "
                f"rows={len(rows)} truncated={truncated} sheets={list(wb.sheetnames)}\n\n"
            )
            out = header + body
            if len(out) > _MAX_OUT:
                out = out[: _MAX_OUT - 20] + "\n…[truncated]"
            return out
        finally:
            wb.close()
    except Exception as exc:  # noqa: BLE001
        return json.dumps({"ok": False, "error": f"excel_read failed: {exc}"}, ensure_ascii=False)


def _sidecar_path(target: Path, kind: str, sheet: str | None = None) -> Path:
    if kind == "xlsx":
        return target.with_name(f"{target.stem}_edited.xlsx")
    safe = re.sub(r"[^\w\-]+", "_", (sheet or "sheet").strip())[:40] or "sheet"
    return target.with_name(f"{target.stem}_{safe}.csv")


def _coerce_cell_value(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, (int, float, bool)):
        return v
    s = str(v)
    if s.lower() in ("", "null", "none"):
        return None
    # keep numbers that look numeric
    try:
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            return float(s)
    except ValueError:
        pass
    return s


def _write_csv_export(path: Path, rows: list[list[Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        for r in rows:
            w.writerow([_cell_str(x) for x in r])


def excel_write(
    cwd: Path,
    path: str,
    *,
    sheet: str | None = None,
    mode: str = "set_cells",
    cells: dict[str, Any] | None = None,
    rows: list[Any] | None = None,
    range: str | None = None,  # noqa: A002
    prefer: str = "inplace",
    trusted_roots: tuple[Path, ...] = (),
) -> str:
    ox = _openpyxl()
    if ox is None:
        return _missing_dep()

    prefer_n = (prefer or "inplace").lower().strip()
    if prefer_n not in ("inplace", "sidecar", "csv"):
        return json.dumps(
            {"ok": False, "error": "prefer must be inplace|sidecar|csv"},
            ensure_ascii=False,
        )
    mode_n = (mode or "set_cells").lower().strip()
    if mode_n not in ("set_cells", "append_rows", "replace_sheet", "export_csv"):
        return json.dumps(
            {
                "ok": False,
                "error": "mode must be set_cells|append_rows|replace_sheet|export_csv",
            },
            ensure_ascii=False,
        )

    try:
        target = file_tools.ensure_under(cwd, file_tools.resolve_path(cwd, path), trusted_roots)
    except PermissionError as exc:
        return json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False)

    bad = _reject_ext(target) if target.suffix else None
    if target.suffix and bad:
        return bad
    if not target.suffix:
        target = target.with_suffix(".xlsx")
    elif target.suffix.lower() != ".xlsx":
        return _reject_ext(target) or json.dumps(
            {"ok": False, "error": "path must end with .xlsx"}, ensure_ascii=False
        )

    # Normalize rows to list[list]
    norm_rows: list[list[Any]] = []
    if rows is not None:
        if not isinstance(rows, list):
            return json.dumps({"ok": False, "error": "rows must be a 2D array"}, ensure_ascii=False)
        for r in rows:
            if isinstance(r, list):
                norm_rows.append(list(r))
            else:
                norm_rows.append([r])

    cells_map = cells if isinstance(cells, dict) else {}

    if mode_n == "set_cells" and not cells_map:
        return json.dumps({"ok": False, "error": "set_cells requires cells object"}, ensure_ascii=False)
    if mode_n in ("append_rows", "replace_sheet") and not norm_rows:
        return json.dumps({"ok": False, "error": f"{mode_n} requires rows"}, ensure_ascii=False)

    # Force CSV export path
    if mode_n == "export_csv" or prefer_n == "csv":
        return _export_csv_flow(
            ox, target, sheet=sheet, range_a1=range, rows_override=norm_rows or None
        )

    try:
        if target.is_file():
            wb = ox.load_workbook(target)
        else:
            wb = ox.Workbook()
            # remove default if we'll rename
            if sheet and wb.active.title != sheet:
                wb.active.title = sheet
    except Exception as exc:  # noqa: BLE001
        # Cannot open — try csv fallback with provided rows/cells
        return _fallback_after_open_fail(
            target, sheet=sheet, mode=mode_n, cells_map=cells_map, norm_rows=norm_rows, err=str(exc)
        )

    try:
        if sheet:
            if sheet in wb.sheetnames:
                ws = wb[sheet]
            else:
                ws = wb.create_sheet(sheet)
        else:
            ws = wb.active

        changed = {"cells": 0, "rows_appended": 0, "rows_written": 0}

        if mode_n == "set_cells":
            for ref, val in cells_map.items():
                col, row = _col_row(str(ref))
                ws.cell(row=row, column=col, value=_coerce_cell_value(val))
                changed["cells"] += 1
        elif mode_n == "append_rows":
            start_row = (ws.max_row or 0) + 1
            if ws.max_row == 1 and all(
                ws.cell(row=1, column=c).value is None for c in range(1, (ws.max_column or 1) + 1)
            ):
                start_row = 1
            for i, row_vals in enumerate(norm_rows):
                for j, val in enumerate(row_vals):
                    ws.cell(row=start_row + i, column=j + 1, value=_coerce_cell_value(val))
                changed["rows_appended"] += 1
        elif mode_n == "replace_sheet":
            # Clear existing cells in used range then write
            if ws.max_row and ws.max_column:
                for r in ws.iter_rows(
                    min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column
                ):
                    for cell in r:
                        cell.value = None
            start_c, start_r = 1, 1
            if range:
                start_c, start_r, _, _ = _parse_range(range.split(":")[0])
            for i, row_vals in enumerate(norm_rows):
                for j, val in enumerate(row_vals):
                    ws.cell(
                        row=start_r + i,
                        column=start_c + j,
                        value=_coerce_cell_value(val),
                    )
                changed["rows_written"] += 1

        out_path = target
        write_mode = "inplace"
        reason = ""

        if prefer_n == "sidecar":
            out_path = _sidecar_path(target, "xlsx")
            write_mode = "sidecar"
            reason = "prefer=sidecar"
            wb.save(out_path)
        else:
            try:
                target.parent.mkdir(parents=True, exist_ok=True)
                wb.save(target)
            except OSError as exc:
                out_path = _sidecar_path(target, "xlsx")
                try:
                    wb.save(out_path)
                    write_mode = "sidecar"
                    reason = f"inplace failed: {exc}"
                except OSError as exc2:
                    return _export_csv_from_ws(
                        ws,
                        _sidecar_path(target, "csv", ws.title),
                        reason=f"xlsx save failed: {exc}; sidecar failed: {exc2}",
                    )

        return json.dumps(
            {
                "ok": True,
                "mode": write_mode,
                "path": str(out_path),
                "sheet": ws.title,
                "reason": reason or None,
                "changed": changed,
            },
            ensure_ascii=False,
        )
    except Exception as exc:  # noqa: BLE001
        return json.dumps({"ok": False, "error": f"excel_write failed: {exc}"}, ensure_ascii=False)
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass


def _export_csv_from_ws(ws, csv_path: Path, *, reason: str) -> str:
    rows, _, _ = _matrix_from_ws(ws, range_a1=None, max_rows=10_000)
    try:
        _write_csv_export(csv_path, rows)
    except OSError as exc:
        return json.dumps(
            {"ok": False, "error": f"csv_export failed: {exc}", "reason": reason},
            ensure_ascii=False,
        )
    return json.dumps(
        {
            "ok": True,
            "mode": "csv_export",
            "path": str(csv_path),
            "sheet": ws.title,
            "reason": reason,
            "changed": {"rows_written": len(rows)},
        },
        ensure_ascii=False,
    )


def _export_csv_flow(ox, target: Path, *, sheet: str | None, range_a1: str | None, rows_override):
    csv_path = _sidecar_path(target, "csv", sheet)
    if rows_override is not None:
        try:
            _write_csv_export(csv_path, rows_override)
            return json.dumps(
                {
                    "ok": True,
                    "mode": "csv_export",
                    "path": str(csv_path),
                    "sheet": sheet or "Sheet1",
                    "reason": "prefer=csv or mode=export_csv",
                    "changed": {"rows_written": len(rows_override)},
                },
                ensure_ascii=False,
            )
        except OSError as exc:
            return json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False)

    if not target.is_file():
        return json.dumps(
            {"ok": False, "error": "export_csv requires existing xlsx or rows"},
            ensure_ascii=False,
        )
    try:
        wb = ox.load_workbook(target, read_only=True, data_only=True)
        try:
            ws = _pick_sheet(wb, sheet)
            rows, _, _ = _matrix_from_ws(ws, range_a1=range_a1, max_rows=10_000)
            _write_csv_export(csv_path, rows)
            return json.dumps(
                {
                    "ok": True,
                    "mode": "csv_export",
                    "path": str(csv_path),
                    "sheet": ws.title,
                    "reason": "mode=export_csv",
                    "changed": {"rows_written": len(rows)},
                },
                ensure_ascii=False,
            )
        finally:
            wb.close()
    except Exception as exc:  # noqa: BLE001
        return json.dumps({"ok": False, "error": f"export_csv failed: {exc}"}, ensure_ascii=False)


def _fallback_after_open_fail(
    target: Path,
    *,
    sheet: str | None,
    mode: str,
    cells_map: dict,
    norm_rows: list,
    err: str,
) -> str:
    # Best-effort: materialize rows from cells or rows into CSV
    rows: list[list[Any]] = list(norm_rows) if norm_rows else []
    if not rows and cells_map:
        # sparse → dense small grid
        coords = []
        for ref, val in cells_map.items():
            try:
                c, r = _col_row(str(ref))
                coords.append((r, c, val))
            except ValueError:
                continue
        if coords:
            max_r = max(r for r, _, _ in coords)
            max_c = max(c for _, c, _ in coords)
            grid = [["" for _ in range(max_c)] for _ in range(max_r)]
            for r, c, val in coords:
                grid[r - 1][c - 1] = _coerce_cell_value(val)
            rows = grid
    if not rows:
        return json.dumps(
            {"ok": False, "error": f"cannot open workbook: {err}; no rows to export"},
            ensure_ascii=False,
        )
    csv_path = _sidecar_path(target, "csv", sheet)
    try:
        _write_csv_export(csv_path, rows)
    except OSError as exc:
        return json.dumps({"ok": False, "error": f"{err}; csv failed: {exc}"}, ensure_ascii=False)
    return json.dumps(
        {
            "ok": True,
            "mode": "csv_export",
            "path": str(csv_path),
            "sheet": sheet or "Sheet1",
            "reason": f"open failed: {err}",
            "changed": {"rows_written": len(rows)},
        },
        ensure_ascii=False,
    )
